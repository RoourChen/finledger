"""CSV 导入导出：表头校验 / 模板 / 编号自动生成 / 金额转换。"""
from db import fetch_one, insert_row
from import_export import (
    _excel_date_to_str,
    _normalize_rate,
    column_label,
    expected_columns,
    grid_to_rows,
    import_rows,
    normalize_rows,
    parse_csv,
    parse_docx,
    parse_excel,
    required_columns,
    split_grid_to_blocks,
    template_csv,
    validate_header,
)


def test_validate_header_missing_and_extra():
    missing, extra = validate_header(["code", "name"], "partners")
    assert missing == ["tax_no", "partner_type", "default_tax_rate", "remark"]
    assert extra == []


def test_validate_header_extra_column():
    full = expected_columns("partners") + ["foo"]
    missing, extra = validate_header(full, "partners")
    assert missing == []
    assert extra == ["foo"]


def test_template_csv_header():
    assert template_csv("partners").strip() == "code,name,tax_no,partner_type,default_tax_rate,remark"


def test_parse_csv():
    text = "code,name,tax_no,partner_type,default_tax_rate,remark\nKH001,甲,TAX1,客户,13%,备注\n"
    fieldnames, rows, missing, extra = parse_csv(text, "partners")
    assert missing == [] and extra == []
    assert len(rows) == 1
    assert rows[0]["name"] == "甲"


def test_parse_csv_detects_missing():
    text = "code,name\nKH001,甲\n"
    fieldnames, rows, missing, extra = parse_csv(text, "partners")
    assert missing == ["tax_no", "partner_type", "default_tax_rate", "remark"]


def test_column_label_code_per_table():
    assert column_label("invoices", "code") == "发票编号（内部）"
    assert column_label("payments", "amount") == "金额"


def test_required_columns():
    assert required_columns("partners") == ["name", "tax_no"]
    assert required_columns("invoices") == ["invoice_number", "partner_code"]


def test_import_rows_applies_defaults(conn):
    # 只给必填列，其余缺省用默认值
    rows = [{"name": "甲", "tax_no": "T1"}]
    ok, errors = import_rows(conn, "partners", rows, "tester")
    assert errors == []
    row = fetch_one(conn, "partners", "code", "KH001")
    assert row["partner_type"] == "客户"
    assert row["default_tax_rate"] == "13%"


def test_import_rows_auto_generates_code(conn):
    rows = [{
        "code": "", "name": "甲客户", "tax_no": "T1",
        "partner_type": "客户", "default_tax_rate": "13%", "remark": "",
    }]
    ok, errors = import_rows(conn, "partners", rows, "tester")
    assert errors == []
    assert ok == ["KH001"]


def test_import_rows_numeric_and_contract_none(conn):
    insert_row(conn, "partners", {
        "code": "KH001", "name": "甲", "tax_no": "T1",
        "partner_type": "客户", "default_tax_rate": "13%", "remark": "",
    }, "tester")
    rows = [{
        "code": "", "invoice_code": "", "invoice_number": "A1", "invoice_date": "2026-01-01",
        "buy_sell_type": "销项", "partner_code": "KH001", "contract_code": "",
        "amount_tax": "113", "amount_no_tax": "", "tax_amount": "",
        "tax_rate": "13%", "status": "正常", "related_invoice_code": "", "remark": "",
    }]
    ok, errors = import_rows(conn, "invoices", rows, "tester")
    assert errors == [], errors
    assert ok == ["FP001"]
    row = conn.execute("SELECT contract_code, amount_tax FROM invoices WHERE code='FP001'").fetchone()
    assert row[0] is None  # 空合同转 NULL
    assert row[1] == 113.0


def test_normalize_rows_chinese_headers():
    fieldnames = ["客户/供应商编号", "名称", "税号", "类型", "默认税率", "备注"]
    rows = [{"客户/供应商编号": "KH001", "名称": "甲", "税号": "T1",
             "类型": "客户", "默认税率": "13%", "备注": ""}]
    norm_fields, norm_rows, unrecognized = normalize_rows("partners", fieldnames, rows)
    assert norm_fields == ["code", "name", "tax_no", "partner_type", "default_tax_rate", "remark"]
    assert norm_rows[0]["name"] == "甲"
    assert unrecognized == []


def test_normalize_rows_unrecognized():
    fieldnames = ["名称", "税号", "乱七八糟"]
    _fields, _rows, unrecognized = normalize_rows("partners", fieldnames, [])
    assert unrecognized == ["乱七八糟"]


def test_parse_excel():
    import io as _io
    import pandas as pd
    buf = _io.BytesIO()
    df = pd.DataFrame({
        "code": ["KH001"], "name": ["甲"], "tax_no": ["T1"],
        "partner_type": ["客户"], "default_tax_rate": ["13%"], "remark": [""],
    })
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="客户", index=False)
    sources = parse_excel(buf.getvalue())
    assert len(sources) == 1
    name, grid = sources[0]
    assert name == "客户"
    fieldnames, rows = grid_to_rows("partners", grid)
    assert rows[0]["name"] == "甲"


def test_parse_excel_numeric_and_date():
    import io as _io
    import pandas as pd
    buf = _io.BytesIO()
    df = pd.DataFrame({
        "amount_tax": [113000.0],
        "invoice_number": [88888888.0],
        "sign_date": [pd.Timestamp("2026-01-01")],
    })
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="S", index=False)
    _name, grid = parse_excel(buf.getvalue())[0]
    # 第一行表头，第二行数据
    assert grid[1][0] == "113000"
    assert grid[1][1] == "88888888"
    assert grid[1][2] == "2026-01-01"


def test_parse_docx():
    import io as _io
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "税号"
    table.cell(1, 0).text = "甲客户"
    table.cell(1, 1).text = "T1"
    buf = _io.BytesIO()
    doc.save(buf)
    sources = parse_docx(buf.getvalue())
    assert len(sources) == 1
    name, grid = sources[0]
    assert name == "表 1"
    _fieldnames, rows = grid_to_rows("partners", grid)
    assert rows[0]["名称"] == "甲客户"


def test_grid_to_rows_skips_title_row():
    grid = [
        ["表1：客户供应商表", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["编号", "名称", "税号", "类型", "默认税率", ""],
        ["KH001", "甲", "T1", "客户", "0.13", ""],
    ]
    fieldnames, rows = grid_to_rows("partners", grid)
    assert "编号" in fieldnames
    assert rows[0]["名称"] == "甲"


def test_normalize_rows_alias_编号():
    fieldnames = ["编号", "名称", "税号"]
    rows = [{"编号": "KH001", "名称": "甲", "税号": "T1"}]
    norm_fields, norm_rows, unrecognized = normalize_rows("partners", fieldnames, rows)
    assert norm_fields == ["code", "name", "tax_no"]
    assert norm_rows[0]["code"] == "KH001"
    assert unrecognized == []


def test_normalize_rate():
    assert _normalize_rate("0.13") == "13%"
    assert _normalize_rate("13") == "13%"
    assert _normalize_rate("13%") == "13%"
    assert _normalize_rate(0.06) == "6%"


def test_import_rows_normalizes_rate(conn):
    rows = [{"name": "甲", "tax_no": "T1", "default_tax_rate": "0.13"}]
    ok, errors = import_rows(conn, "partners", rows, "tester")
    assert errors == []
    row = fetch_one(conn, "partners", "code", "KH001")
    assert row["default_tax_rate"] == "13%"


def test_split_grid_to_blocks():
    grid = [
        ["表1：客户供应商表", "", ""],
        ["", "", ""],
        ["编号", "名称", "税号"],
        ["KH001", "甲", "T1"],
        ["表2：合同表", "", ""],
        ["合同编号", "名称", "金额"],
        ["HT001", "合同A", "100"],
    ]
    blocks = split_grid_to_blocks(grid)
    assert len(blocks) == 2
    assert blocks[0][0] == "表1：客户供应商表"
    assert blocks[1][0] == "表2：合同表"
    assert len(blocks[0][1]) == 3  # 空行 + 表头 + 1 数据


def test_excel_date_to_str():
    assert _excel_date_to_str(45301) == "2024-01-10"
    assert _excel_date_to_str("2024-01-10") == "2024-01-10"
    assert _excel_date_to_str("") == ""
    assert _excel_date_to_str("不是日期") == "不是日期"


def test_import_rows_converts_excel_date(conn):
    insert_row(conn, "partners", {
        "code": "KH001", "name": "甲", "tax_no": "T1",
        "partner_type": "客户", "default_tax_rate": "13%", "remark": "",
    }, "tester")
    rows = [{
        "partner_code": "KH001", "name": "合同A",
        "sign_date": "45301", "amount_tax": "113000", "amount_no_tax": "100000",
    }]
    ok, errors = import_rows(conn, "contracts", rows, "tester")
    assert errors == []
    row = fetch_one(conn, "contracts", "code", "HT001")
    assert row["sign_date"] == "2024-01-10"


def test_import_rows_reports_numeric_error(conn):
    insert_row(conn, "partners", {
        "code": "KH001", "name": "甲", "tax_no": "T1",
        "partner_type": "客户", "default_tax_rate": "13%", "remark": "",
    }, "tester")
    rows = [{
        "code": "", "invoice_code": "", "invoice_number": "A1", "invoice_date": "",
        "buy_sell_type": "销项", "partner_code": "KH001", "contract_code": "",
        "amount_tax": "abc", "amount_no_tax": "", "tax_amount": "",
        "tax_rate": "13%", "status": "正常", "related_invoice_code": "", "remark": "",
    }]
    ok, errors = import_rows(conn, "invoices", rows, "tester")
    assert ok == []
    assert any("不是数字" in e for e in errors)
